"""Word 报告总览表（Overview）与段落/表格替换

从 word_report_writer.py 拆出（大文件拆分重构）。"""

import logging
import math

from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from battery_analysis.reports import word_utils
from battery_analysis.utils.report_coordinator import compute_report_content_base

logger = logging.getLogger(__name__)


class WordOverviewMixin:
    """总览内容计算、总览表写入、占位符替换与表格插入"""

    # ── 计算Overview表内容 ──

    def _prepare_overview_content(self, doc, stats):
        """计算Overview表所需的位置参数和内容列表"""
        base = compute_report_content_base(
            self.listCurrentLevel,
            self.intCurrentLevelNum,
            self.listVoltageLevel,
            self.intVoltageLevelNum,
            self.listTestInfo,
            self.listBatteryInfo,
            self.strSampleXlsxPath,
            self.strResultXlsxPath,
            self.strReportWordPath,
            stats,
        )

        intTestDateStartRow = base["intTestProfileStartLine"] + 11

        # Word 特有的 listTestInfoForReplace 副作用
        if base["strResult"] == "Pass":
            self.listTestInfoForReplace[8] = "meets"
            self.listTestInfoForReplace[10] = "Pass"
        elif base["strResult"] == "Incomplete":
            # P3 第三态：无数据格 → 不足以下 Pass/Fail 结论（产线需复测）
            self.listTestInfoForReplace[8] = "has not been verified against"
            self.listTestInfoForReplace[10] = "Incomplete"
        else:
            self.listTestInfoForReplace[8] = "doesn't meet"
            self.listTestInfoForReplace[10] = "Fail"
        self.listTestInfoForReplace[9] = base["strRequiredUseableCapacityPercentage"]
        self.listTestInfoForReplace[7] = (
            f"{math.floor(100 * stats['mm2s'][base['intPosiMaxmA']][base['intPosi2V25']] / int(self.listTestInfo[9]))}%"
        )

        return {
            **base,
            "intTestDateStartRow": intTestDateStartRow,
        }

    # ── 写入Overview表（Word） ──

    def _write_overview_table(self, doc, content, stats):
        """在Word文档中写入Overview表"""
        intPosiMaxmA = content["intPosiMaxmA"]
        intPosi2V25 = content["intPosi2V25"]
        intTestProfileStartLine = content["intTestProfileStartLine"]
        intTestDateStartRow = content["intTestDateStartRow"]
        intActualMeasuredCapacityLength = content["intActualMeasuredCapacityLength"]
        listStrItems = content["listStrItems"]
        listStrContent = content["listStrContent"]

        table = doc.add_table(
            intTestDateStartRow + 6, 1 + self.intVoltageLevelNum * 2, style="Table Grid"
        )
        for row in range(intTestDateStartRow + 6):
            table.rows[row].height = Cm(0.5)
            for col in range(1 + self.intVoltageLevelNum * 2):
                cell = table.cell(row, col)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                cell.paragraphs[0].paragraph_format.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell.paragraphs[0].paragraph_format.line_spacing_rules = WD_LINE_SPACING.SINGLE
                cell.paragraphs[0].paragraph_format.space_after = Pt(0)
                if row == intTestDateStartRow + 3:
                    word_utils.table_set_bg_color(cell, "#FFFF00")

        # merge cells
        for row in range(intTestDateStartRow + 6):
            if row <= intTestProfileStartLine + 6 or row >= intTestDateStartRow:
                cell1 = table.cell(row, 1)
                for col in range(1, self.intVoltageLevelNum * 2):
                    cell2 = table.cell(row, col + 1)
                    cell1.merge(cell2)
            else:
                if row == intTestProfileStartLine + 7:
                    cell1 = table.cell(row, 1)
                    for col in range(1, self.intVoltageLevelNum):
                        cell2 = table.cell(row, col + 1)
                        cell1.merge(cell2)
                    cell1 = table.cell(row, 1 + self.intVoltageLevelNum)
                    for col in range(1 + self.intVoltageLevelNum, self.intVoltageLevelNum * 2):
                        cell2 = table.cell(row, col + 1)
                        cell1.merge(cell2)
                else:
                    for v in range(self.intVoltageLevelNum):
                        cell1 = table.cell(row, 1 + v * 2)
                        cell2 = table.cell(row, 2 + v * 2)
                        cell1.merge(cell2)

        cell1 = table.cell(intTestProfileStartLine + 8, 0)
        cell2 = table.cell(intTestProfileStartLine + 9, 0)
        cell1.merge(cell2)
        cell2 = table.cell(intTestProfileStartLine + 10, 0)
        cell1.merge(cell2)

        # 写入标签列
        for i in range(4):
            table.cell(i, 0).paragraphs[0].add_run(listStrItems[i])
        if intTestProfileStartLine == 3:
            # 构造方法为空：第 3 行标签无意义，清除后由 Test Profile 占用
            table.cell(3, 0).paragraphs[0].text = (
                table.cell(3, 0).paragraphs[0].text.replace(listStrItems[3], "")
            )
        for i in range(4, 12):
            table.cell(intTestProfileStartLine + (i - 4), 0).paragraphs[0].add_run(listStrItems[i])
        table.cell(intTestProfileStartLine + 8, 0).paragraphs[0].add_run(listStrItems[13])
        for i in range(14, 20):
            table.cell(intTestDateStartRow + (i - 14), 0).paragraphs[0].add_run(listStrItems[i])

        # 写入内容列
        for i in range(4):
            table.cell(i, 1).paragraphs[0].add_run(listStrContent[i])
        if intTestProfileStartLine == 3:
            # 构造方法为空：清除第 3 行内容占位
            table.cell(3, 1).paragraphs[0].text = (
                table.cell(3, 1).paragraphs[0].text.replace(listStrContent[3], "")
            )
        table.cell(intTestProfileStartLine, 1).paragraphs[0].text = ""
        if len(listStrContent[4].split("\\")) == 1:
            table.cell(intTestProfileStartLine, 1).paragraphs[0].add_run(listStrContent[4])
        else:
            word_utils.add_hyperlink(
                table.cell(intTestProfileStartLine, 1).paragraphs[0],
                listStrContent[4],
                listStrContent[4].split("\\")[-1],
            )
        for i in range(5, 12):
            table.cell(intTestProfileStartLine + (i - 4), 1).paragraphs[0].add_run(
                listStrContent[i]
            )
        table.cell(intTestProfileStartLine + 7, 1 + self.intVoltageLevelNum).paragraphs[0].add_run(
            listStrContent[12]
        )

        counts = stats.get(
            "counts", [[1] * self.intVoltageLevelNum] * self.intCurrentLevelNum
        )
        for v in range(self.intVoltageLevelNum):
            text1 = (
                table.cell(intTestProfileStartLine + 8, 1 + v * 2)
                .paragraphs[0]
                .add_run(f"{self.listVoltageLevel[v]}V")
            )
            text2 = (
                table.cell(intTestProfileStartLine + 9, 1 + v * 2)
                .paragraphs[0]
                .add_run(
                    "No data (n=0)" if not counts[intPosiMaxmA][v]
                    else f"{round(stats['mm2s'][intPosiMaxmA][v], 2)}"
                )
            )
            text3 = (
                table.cell(intTestProfileStartLine + 10, 1 + v * 2)
                .paragraphs[0]
                .add_run(
                    ""
                    if not counts[intPosiMaxmA][v]
                    else f"{math.floor(100 * stats['mm2s'][intPosiMaxmA][v] / int(listStrContent[10]))}%"
                )
            )
            if v == intPosi2V25:
                text1.font.bold = True
                text2.font.bold = True
                text3.font.bold = True

        for i in range(14, 20):
            if i == 18:
                table.cell(intTestDateStartRow + (i - 14), 1).paragraphs[0].text = ""
                word_utils.add_hyperlink(
                    table.cell(intTestDateStartRow + (i - 14), 1).paragraphs[0],
                    listStrContent[i],
                    listStrContent[i].split("\\")[-1],
                )
            else:
                table.cell(intTestDateStartRow + (i - 14), 1).paragraphs[0].add_run(
                    listStrContent[i]
                )

        for row in range(intTestDateStartRow + 6):
            for col in range(1 + self.intVoltageLevelNum * 2):
                cell = table.cell(row, col)
                runs = cell.paragraphs[0].runs
                for run in runs:
                    run.font.size = Pt(9)
        table.cell(0, 0).width = Cm(27)

        return table

    # ── 文本替换、图片插入、表格插入 ──

    def _replace_and_insert(self, doc, tables):
        """替换占位符文本、插入图片，并在指定位置插入表格"""
        tableOverview = tables["overview"]
        tableVersionHistory = tables["version_history"]
        tableTestInformation = tables["test_information"]
        tableStatisticalsResults = tables["statistical_results"]

        bInsertOverview = False
        bInsertVersionHistory = False
        bInsertTestInformation = False
        bInsertStatisticalsResults = False
        intStepOut = 0

        for paragraph in doc.paragraphs:
            modified = False

            # 替换文本
            for t in range(len(self.listTextToReplace)):
                if self.listTextToReplace[t] in paragraph.text:
                    modified = True
                    if self.listTextToReplace[t] == "StrD":
                        paragraph.text = paragraph.text.replace(self.listTextToReplace[t], "")
                        text = paragraph.add_run(f"{self.listTestInfoForReplace[t]}")
                        text.font.bold = True
                        paragraph.add_run(".")
                    else:
                        paragraph.text = paragraph.text.replace(
                            self.listTextToReplace[t], f"{self.listTestInfoForReplace[t]}"
                        )

            # 替换图片
            if not modified:
                for i in range(len(self.listImageToReplace)):
                    if self.listImageToReplace[i] in paragraph.text:
                        paragraph.text = paragraph.text.replace(self.listImageToReplace[i], "")
                        if i < 2 * len(self.listPngPath) + 1:
                            if i == 0:
                                paragraph.add_run("").add_picture(
                                    self.strFilteredPngPath, width=Cm(15)
                                )
                            elif i % 2 == 1:
                                paragraph.add_run("").add_picture(
                                    self.listPngPath[int((i - 1) / 2)], width=Cm(7.2)
                                )
                            else:
                                paragraph.add_run(
                                    f"Figure {int(i / 2 + 1)}  {self.listTestInfoForReplace[2]} "
                                    f"{self.listTestInfoForReplace[0]}-{self.listTestInfoForReplace[1]} "
                                    f"Boxplot, {self.listCurrentLevel[int(i / 2 - 1)]}mA"
                                )
                        else:
                            paragraph._element.getparent().remove(paragraph._element)
                        continue

            # 识别插入点并插入表格
            if (
                "Battery Quality Test / Alternative Battery Test for ESL Batteries"
                in paragraph.text
            ):
                bInsertOverview = True
                intStepOut = 1
            elif "Version history" in paragraph.text and paragraph.style.name == "Heading 2":
                bInsertVersionHistory = True
                intStepOut = 0
            elif "Test Information" in paragraph.text and paragraph.style.name == "Heading 1":
                bInsertTestInformation = True
                intStepOut = 0
            elif "Test results" in paragraph.text and paragraph.style.name == "Heading 1":
                bInsertStatisticalsResults = True
                intStepOut = 2

            if intStepOut:
                intStepOut = intStepOut - 1
            else:
                if bInsertOverview:
                    bInsertOverview = False
                    paragraph._p.addnext(tableOverview._tbl)
                    # 保留表格后的一个空段落，不删除包含分页符的段落。
                    following = tableOverview._tbl.getnext()
                    blank_count = 0
                    while following is not None and following.tag == qn("w:p"):
                        if following.xpath(".//w:t | .//w:br | .//w:drawing | .//w:sectPr"):
                            break
                        next_element = following.getnext()
                        blank_count += 1
                        if blank_count > 1:
                            following.getparent().remove(following)
                        following = next_element
                elif bInsertVersionHistory:
                    bInsertVersionHistory = False
                    paragraph._p.addnext(tableVersionHistory._tbl)
                elif bInsertTestInformation:
                    bInsertTestInformation = False
                    paragraph._p.addnext(tableTestInformation._tbl)
                elif bInsertStatisticalsResults:
                    bInsertStatisticalsResults = False
                    paragraph._p.addnext(tableStatisticalsResults._tbl)

            # 删除温度符号
            if self.listTestInfo[7] == "Room Temperature" and "℃" in paragraph.text:
                paragraph.text = paragraph.text.replace("℃", "")
