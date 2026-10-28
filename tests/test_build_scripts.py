"""测试 scripts/build.py 的构建入口与失败检测逻辑"""

import runpy
import subprocess
import sys
import tomllib
from pathlib import Path
from textwrap import dedent
from types import ModuleType
from unittest.mock import Mock

import pytest

from scripts.build import _build_failed


@pytest.mark.parametrize("upx_available", [False, True])
@pytest.mark.parametrize("build_type", ["Debug", "Release"])
def test_pyinstaller_entry_matches_project_script(
    monkeypatch, tmp_path, build_type, upx_available
):
    from scripts.build import BuildManager

    # 只检查配置与命令，不清理真实产物或调用 PyInstaller。
    monkeypatch.setattr(BuildManager, "clean_build_dirs", lambda self: None)
    monkeypatch.setattr("scripts.build.shutil.which", lambda name: "upx" if upx_available else None)
    manager = BuildManager(build_type)
    project = tomllib.loads(
        (manager.project_root / "pyproject.toml").read_text(encoding="utf-8")
    )
    entry_module = project["project"]["scripts"]["battery-analysis"].split(":")[0]
    expected_entry = (
        manager.project_root / "src" / entry_module.replace(".", "/")
    ).with_suffix(".py")
    app_config = manager.apps_config[0]
    app_config["exe_name"] = "test-battery-analyzer"

    args = manager._build_pyinstaller_args(
        app_config, tmp_path, manager.project_root / "src", tmp_path / "dist"
    )

    assert expected_entry.is_file()
    assert app_config["main_file_path"] == expected_entry
    assert str(expected_entry) in args
    assert "--strip" not in args
    assert "--noupx" in args
    assert "--exclude-module=PIL.AvifImagePlugin" in args
    assert "--exclude-module=lxml.objectify" in args
    assert "--exclude-module=lxml.etree" not in args
    hooks_dir = manager.project_root / "scripts" / "pyinstaller_hooks"
    assert f"--additional-hooks-dir={hooks_dir}" in args
    assert (hooks_dir / "hook-PyQt6.QtGui.py").is_file()
    import numpy

    numpy_libs = Path(numpy.__file__).parent.parent / "numpy.libs"
    assert numpy_libs.is_dir()
    assert f"--add-binary={numpy_libs};numpy.libs" in args


@pytest.mark.parametrize("pdf_plugin", [None, "qpdf.dll", "QPDF.DLL"])
def test_qtgui_hook_only_removes_pdf_reader(monkeypatch, tmp_path, pdf_plugin):
    hook = Path(__file__).resolve().parents[1] / "scripts/pyinstaller_hooks/hook-PyQt6.QtGui.py"
    retained = [
        (str(tmp_path / name), "PyQt6/Qt6")
        for name in ("qsvg.dll", "qjpeg.dll", "opengl32sw.dll", "Qt6Network.dll")
    ]
    binaries = retained.copy()
    if pdf_plugin is not None:
        binaries.insert(1, (str(tmp_path / pdf_plugin), "PyQt6/Qt6/plugins/imageformats"))
    hiddenimports = ["PyQt6.QtCore"]
    datas = [("qtbase_zh_CN.qm", "PyQt6/Qt6/translations")]
    collect_dependencies = Mock(return_value=(hiddenimports, binaries, datas))
    # 不要求开发测试环境安装 PyInstaller，钩子只依赖此官方收集接口。
    qt_hooks = ModuleType("PyInstaller.utils.hooks.qt")
    qt_hooks.add_qt6_dependencies = collect_dependencies
    monkeypatch.setitem(sys.modules, qt_hooks.__name__, qt_hooks)

    result = runpy.run_path(str(hook))

    collect_dependencies.assert_called_once_with(str(hook))
    assert result["binaries"] == retained
    assert result["hiddenimports"] is hiddenimports
    assert result["datas"] is datas


def test_image_exports_and_office_reports_work_without_unused_modules(tmp_path):
    # 独立进程避免其他测试预先加载已排除模块，模拟冻结包中的依赖边界。
    result = subprocess.run(
        [sys.executable, "-c", dedent("""
            import sys
            from pathlib import Path
            from zipfile import ZipFile

            sys.modules["PIL.AvifImagePlugin"] = None
            sys.modules["PIL._avif"] = None
            sys.modules["lxml.objectify"] = None

            from docx import Document
            from matplotlib.backends.backend_agg import FigureCanvasAgg
            from matplotlib.figure import Figure
            from openpyxl import Workbook
            from openpyxl.drawing.image import Image as ExcelImage
            from PIL import Image
            from xlsxwriter import Workbook as XlsxWriterWorkbook

            Image.init()
            assert "AVIF" not in Image.OPEN
            assert "AVIF" not in Image.SAVE
            figure = Figure(figsize=(1, 1))
            canvas = FigureCanvasAgg(figure)
            figure.add_subplot().plot([0, 1], [0, 1])
            formats = canvas.get_supported_filetypes()
            assert "avif" not in formats
            for fmt in formats:
                if fmt == "pgf":  # PGF 依赖外部 LaTeX，与 Pillow 无关。
                    continue
                output = Path(f"chart.{fmt}")
                figure.savefig(output, format=fmt)
                assert output.stat().st_size > 0, fmt

            document = Document()
            document.add_picture("chart.png")
            document.save("report.docx")
            assert len(Document("report.docx").inline_shapes) == 1
            templates = list(Path(sys.argv[1]).glob("*.docx"))
            assert len(templates) == 2
            for index, template in enumerate(templates):
                document = Document(template)
                original_images = len(document.inline_shapes)
                document.add_picture("chart.png")
                output = f"template_report_{index}.docx"
                document.save(output)
                assert len(Document(output).inline_shapes) == original_images + 1

            workbook = Workbook()
            workbook.active.add_image(ExcelImage("chart.png"), "A1")
            workbook.save("openpyxl.xlsx")
            with XlsxWriterWorkbook("xlsxwriter.xlsx") as workbook:
                workbook.add_worksheet().insert_image("A1", "chart.png")
            for filename in ("openpyxl.xlsx", "xlsxwriter.xlsx"):
                with ZipFile(filename) as archive:
                    assert "xl/media/image1.png" in archive.namelist()
        """), str(Path(__file__).resolve().parents[1] / "src/battery_analysis/templates")],
        cwd=tmp_path, capture_output=True, text=True, timeout=60,
    )
    assert result.returncode == 0, result.stdout + result.stderr


class TestBuildFailed:
    def test_returncode_zero_is_success(self):
        assert _build_failed(Mock(returncode=0)) is False

    def test_returncode_nonzero_is_failure(self):
        assert _build_failed(Mock(returncode=1)) is True
        assert _build_failed(Mock(returncode=-1)) is True

    def test_none_result_is_failure(self):
        # _execute_pyinstaller_command 异常时返回 CompletedProcess(_, 1)，不可能为 None，
        # 但防御性处理 None 为失败
        assert _build_failed(None) is True


class TestBuildFailurePropagation:
    """集成回归：PyInstaller 命令非零时 build() 必须 SystemExit(1)，堵住 CI 假绿。"""

    def _make_manager(self, monkeypatch, tmp_path):
        import subprocess

        from scripts.build import BuildManager

        # 避免 BuildManager.__init__ 清理/创建真实构建目录
        monkeypatch.setattr(BuildManager, "clean_build_dirs", lambda self: None)
        manager = BuildManager("Release")
        manager.project_root = tmp_path
        manager.temp_build_dir = tmp_path / "work"
        manager.build_path = manager.temp_build_dir
        return manager, subprocess

    def test_pyinstaller_nonzero_exits(self, monkeypatch, tmp_path):
        manager, subprocess = self._make_manager(monkeypatch, tmp_path)
        # build() 内部直接调用 subprocess.run，需 patch 到 subprocess 模块
        monkeypatch.setattr(
            subprocess,
            "run",
            lambda *a, **kw: subprocess.CompletedProcess(a[0], 1),
        )
        with pytest.raises(SystemExit) as excinfo:
            manager.build()
        assert excinfo.value.code == 1

    def test_pyinstaller_zero_does_not_exit(self, monkeypatch, tmp_path):
        manager, subprocess = self._make_manager(monkeypatch, tmp_path)
        monkeypatch.setattr(
            subprocess,
            "run",
            lambda *a, **kw: subprocess.CompletedProcess(a[0], 0),
        )
        manager.build()  # 不抛 SystemExit 即通过

    def test_missing_artifact_propagates_and_preserves_work(self, monkeypatch, tmp_path):
        manager, subprocess = self._make_manager(monkeypatch, tmp_path)
        monkeypatch.setattr(
            subprocess, "run", lambda *a, **kw: subprocess.CompletedProcess(a[0], 0)
        )
        with pytest.raises(FileNotFoundError, match="缺少构建产物"):
            manager.run_build()
        assert manager.temp_build_dir.is_dir()

    def test_io_error_propagates(self, monkeypatch, tmp_path):
        manager, _ = self._make_manager(monkeypatch, tmp_path)
        monkeypatch.setattr(manager, "build", Mock(side_effect=PermissionError("locked")))
        with pytest.raises(PermissionError, match="locked"):
            manager.run_build()

    @pytest.mark.parametrize("no_clean", [False, True])
    def test_work_cleanup_after_verified_build(self, monkeypatch, tmp_path, no_clean):
        manager, subprocess = self._make_manager(monkeypatch, tmp_path)
        manager.no_clean = no_clean
        monkeypatch.setenv("CI", "true")

        def fake_build(*args, **kwargs):
            output = tmp_path / "build" / "Release"
            exe_name = manager.apps_config[0]["exe_name"]
            (output / f"{exe_name}.exe").write_bytes(b"test exe")
            (manager.temp_build_dir / "Analysis-00.toc").write_text("cache")
            return subprocess.CompletedProcess(args[0], 0)

        monkeypatch.setattr(subprocess, "run", fake_build)
        manager.run_build()
        assert manager.temp_build_dir.exists() is no_clean


def test_window_icon_is_bundled_at_runtime_path(monkeypatch, tmp_path):
    from scripts.build import BuildManager

    monkeypatch.setattr(BuildManager, "clean_build_dirs", lambda self: None)
    manager = BuildManager("Release")
    config = manager.apps_config[0]
    config["exe_name"] = "test"
    args = manager._build_pyinstaller_args(
        config, tmp_path, manager.project_root / "src", tmp_path / "dist"
    )
    icon = manager.project_root / "src" / "battery_analysis" / "resources" / "icons" / config["icon_name"]
    assert icon.is_file()
    assert f"--icon={icon}" in args
    assert f"{icon};battery_analysis/resources/icons" in args
