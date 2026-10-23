"""Cover-page overview placement in both Word report templates."""

from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from battery_analysis.reports.word_overview import WordOverviewMixin


@pytest.mark.parametrize("suffix", ["", "_MP"])
def test_cover_overview_has_one_blank_paragraph_before_and_after(suffix, tmp_path):
    template_dir = Path(__file__).resolve().parents[3] / "src/battery_analysis/templates"
    doc = Document(template_dir / f"Battery Measurement Report of TypeC TypeA_TypeD{suffix}.docx")
    writer = WordOverviewMixin()
    writer.listTextToReplace = []
    writer.listImageToReplace = []
    writer.listTestInfo = [""] * 8
    tables = {
        name: doc.add_table(rows=1, cols=1)
        for name in ("overview", "version_history", "test_information", "statistical_results")
    }
    tables["overview"].cell(0, 0).text = "Overview table"

    writer._replace_and_insert(doc, tables)
    output = tmp_path / "report.docx"
    doc.save(output)
    result = Document(output)

    title = next(
        p for p in result.paragraphs
        if "Battery Quality Test / Alternative Battery Test for ESL Batteries" in p.text
    )
    spacer = title._p.getnext()
    assert spacer.tag == qn("w:p")
    assert not spacer.xpath(".//w:t | .//w:br | .//w:drawing")
    overview = spacer.getnext()
    assert overview.tag == qn("w:tbl")
    assert overview.xpath(".//w:t")[0].text == "Overview table"
    after_table = overview.getnext()
    assert after_table.tag == qn("w:p")
    assert not after_table.xpath(".//w:t | .//w:br | .//w:drawing")
    page_break = after_table.getnext()
    assert page_break.tag == qn("w:p")
    assert page_break.xpath('.//w:br[@w:type="page"]')
    assert any(p.text == "Table of contents" for p in result.paragraphs)
