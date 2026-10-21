"""新增电池型号端到端链路测试

场景：用户在 Rules 配置中新增一个从未见过的电池型号（如 XYZ1430、CR2032），
填入容量等参数规则后，完整链路必须正常工作：

  规则录入 → derive_specifications（规格下拉自动派生）
           → UI 参数填充（85%|75% 构造相关换算 / 绝对值）
           → 脉冲电流查表（§7.1.1 容量分档）
           → build_expectations（期望可用容量格派生）
           → compute_statistics（µ/µ-2σ/最小值）
           → compute_report_content_base（§6.2.11 Pass/Fail + Remarks）
           → ExcelReportWriter.write（Overview 判定标注落盘）

同时覆盖两个安全网：
  - 期望格不在实测档位 → 回退历史单点判定（行为与旧版一致）
  - Coin/Pouch 两类新前缀型号的分类兜底（非 CR/CP/CF 前缀按容量分档）
"""

from types import SimpleNamespace

from openpyxl import load_workbook

from battery_analysis.main.business_logic.validation_spec_rules import (
    ValidationSpecRulesMixin,
)
from battery_analysis.utils.battery_classifier import derive_specifications
from battery_analysis.utils.pulse_currents import lookup_pulse_currents
from battery_analysis.utils.quality_requirements import build_expectations
from battery_analysis.utils.report_coordinator import compute_report_content_base
from battery_analysis.utils.writers.excel_report_writer import ExcelReportWriter
from battery_analysis.utils.writers.statistics_utils import (
    compute_list_cpt,
    compute_statistics,
)

# 用户新录入的规则（保留默认 CR2450，新增两个从未见过的型号）
NEW_RULES = [
    "CR2450/1S1P/600/550/380/1.0",
    "XYZ1430/1S1P/1500/1500/85%|75%/5",  # Pouch，非 CR/CP/CF 前缀（容量兜底分类）
    "CR2032/1S1P/220/200/200/1.0",  # Coin，绝对值期望
]

CURRENTS = [60, 30, 15, 6]  # XYZ1430 (1500mAh → 1200-2000 档) 的查表结果
VOLTAGES = [2.6, 2.5, 2.4, 2.3, 2.25, 2.2, 2.1, 2.0, 1.8]  # 9 档超集
N_C = len(CURRENTS)
N_V = len(VOLTAGES)
TARGET_C = 0  # 60mA
TARGET_V = 4  # 2.25V


# ── helpers ──


def make_test_info(spec, construction, method, qnom, required, currents, voltages):
    """按 TestInfo to_list 的 19 位顺序构造（idx2=规格、idx9=计算容量、idx17=要求可用容量）"""
    return [
        "Pouch Cell" if construction else "Coin Cell",  # 0 battery_type
        construction,  # 1 construction_method
        spec,  # 2 specification_type
        method,  # 3 specification_method
        "Ewin",  # 4 manufacturer
        "DC2401",  # 5 batch_date_code
        "3",  # 6 samples_qty
        "25",  # 7 temperature_value
        str(qnom),  # 8 datasheet_nominal_capacity
        str(qnom),  # 9 calculation_nominal_capacity
        "0",  # 10 accelerated_aging
        "Lab-A",  # 11 tester_location
        "Tester",  # 12 tested_by
        "/path/to/test.profile",  # 13 test_profile
        list(currents),  # 14 current_levels
        list(voltages),  # 15 voltage_levels
        "1.0",  # 16 version
        str(required),  # 17 required_usable_capacity
        "Reviewer",  # 18 reported_by
    ]


def make_charge(target_values):
    """3 只电池 × C*V 平铺容量（idx = c*N_V + v），默认 1500，目标格用给定值"""
    charge = []
    for b in range(3):
        row = [1500] * (N_C * N_V)
        row[TARGET_C * N_V + TARGET_V] = target_values[b]
        charge.append(row)
    return charge


def make_battery_info(charge):
    return [
        charge,
        ["BAT-001", "BAT-002", "BAT-003"],
        ["2024-01-15 10:00:00", "2024-01-15 14:30:00"],
    ]


def run_report(test_info, battery_info):
    stats = compute_statistics(compute_list_cpt(battery_info[0], 3, N_C, N_V), N_C, N_V)
    return compute_report_content_base(
        test_info[14],
        N_C,
        test_info[15],
        N_V,
        test_info,
        battery_info,
        "/tmp/out/sample.xlsx",
        "/tmp/out/result.xlsx",
        "/tmp/out/report.docx",
        stats,
    )


# ── 1. 规则录入 → 规格派生（下拉框数据源）──


class TestSpecDerivation:
    def test_new_pouch_spec_appears_in_derived_list(self):
        specs = derive_specifications(NEW_RULES)
        assert "XYZ1430" in specs["Pouch Cell"]
        assert "CR2032" in specs["Coin Cell"]
        assert "CR2450" in specs["Coin Cell"]
        assert "XYZ1430" not in specs["Coin Cell"]

    def test_derived_lists_match_dropdown_source(self):
        """config_manager.get_config("SpecificationType*") 的派生与此同源"""
        specs = derive_specifications(NEW_RULES)
        all_specs = specs["Coin Cell"] + specs["Pouch Cell"]
        assert "XYZ1430" in all_specs  # _get_all_specifications = coin + pouch


# ── 2. UI 参数填充（规则第 5 段三种形式）──


class FakeLineEdit:
    def __init__(self):
        self.value = None

    def setText(self, text):
        self.value = text


class FakeCombo:
    def __init__(self, text=""):
        self._text = text

    def currentText(self):
        return self._text


class FakeConfigManager:
    def __init__(self, rules):
        self._rules = rules

    def get_config(self, key):
        if key == "BatteryConfig/Rules":
            return self._rules
        return None


def make_mixin(spec, method, construction, rules=NEW_RULES):
    mixin = ValidationSpecRulesMixin()
    mixin.main_window = SimpleNamespace(
        lineEdit_DatasheetNominalCapacity=FakeLineEdit(),
        lineEdit_CalculationNominalCapacity=FakeLineEdit(),
        lineEdit_RequiredUseableCapacity=FakeLineEdit(),
        comboBox_Specification_Type=FakeCombo(spec),
        comboBox_Specification_Method=FakeCombo(method),
        comboBox_ConstructionMethod=FakeCombo(construction),
        config_manager=FakeConfigManager(rules),
    )
    return mixin


class TestRuleFill:
    def test_new_pouch_rule_spiral(self):
        mixin = make_mixin("XYZ1430", "1S1P", "Spiral Type")
        mixin._apply_specification_rules()
        assert mixin.main_window.lineEdit_DatasheetNominalCapacity.value == "1500"
        assert mixin.main_window.lineEdit_CalculationNominalCapacity.value == "1500"
        assert mixin.main_window.lineEdit_RequiredUseableCapacity.value == "1275"  # 85%

    def test_new_pouch_rule_laminate(self):
        mixin = make_mixin("XYZ1430", "1S1P", "Laminate Type")
        mixin._apply_specification_rules()
        assert mixin.main_window.lineEdit_RequiredUseableCapacity.value == "1125"  # 75%

    def test_check_construction_switches_value(self):
        mixin = make_mixin("XYZ1430", "1S1P", "Spiral Type")
        mixin.check_construction()
        assert mixin.main_window.lineEdit_RequiredUseableCapacity.value == "1275"
        mixin.main_window.comboBox_ConstructionMethod = FakeCombo("Laminate Type")
        mixin.check_construction()
        assert mixin.main_window.lineEdit_RequiredUseableCapacity.value == "1125"

    def test_new_coin_rule_absolute_value(self):
        mixin = make_mixin("CR2032", "1S1P", "")
        mixin._apply_specification_rules()
        assert mixin.main_window.lineEdit_RequiredUseableCapacity.value == "200"

    def test_method_mismatch_clears_fields(self):
        mixin = make_mixin("XYZ1430", "1S2P", "Spiral Type")  # 规则是 1S1P
        mixin._apply_specification_rules()
        assert mixin.main_window.lineEdit_RequiredUseableCapacity.value == ""  # 方法不匹配 → 清空


# ── 3. 脉冲电流查表（新型号走容量分档兜底）──


class TestPulseLookup:
    def test_new_pouch_model_band_lookup(self):
        # 1500mAh → 1200≤Q<2000 档（§7.1.1.5/.6）
        assert lookup_pulse_currents("XYZ1430", 1500, "1S1P") == [60, 30, 15, 6]

    def test_new_coin_model_standard_profile(self):
        assert lookup_pulse_currents("CR2032", 220, "1S1P") == [30, 26, 15, 6]

    def test_new_cp_prefixed_small_pouch(self):
        # CP 前缀 → Pouch；500mAh < 1200 → 30/26/15/6
        assert lookup_pulse_currents("CP080512", 500, "1S1P") == [30, 26, 15, 6]

    def test_unknown_spec_empty_capacity_returns_no_override(self):
        # CP 前缀（Pouch）+ 容量未知 → 不给档位（不覆盖已有配置）
        assert lookup_pulse_currents("CP080512", 0, "1S1P") == []

    def test_unknown_prefix_zero_capacity_falls_back_to_coin(self):
        # 非 CR/CP/CF 前缀 + 容量 0 → classify 兜底 Coin → 标准档
        # （实际链路中容量来自规则/文件名，正常不会为 0）
        assert lookup_pulse_currents("XYZ1430", 0, "1S1P") == [30, 26, 15, 6]


# ── 4. 期望格派生（§7.1.1 → build_expectations）──


class TestExpectations:
    def test_new_pouch_expectation_at_pdf_cell(self):
        ti = make_test_info("XYZ1430", "Spiral Type", "1S1P", 1500, 1275, CURRENTS, VOLTAGES)
        exps = build_expectations(ti)
        assert len(exps) == 1
        e = exps[0]
        assert (e.current_ma, e.cutoff_v, e.threshold_mah) == (60.0, 2.25, 1275)

    def test_new_coin_expectation_at_30ma(self):
        ti = make_test_info("CR2032", "", "1S1P", 220, 200, [30, 26, 15, 6], VOLTAGES)
        exps = build_expectations(ti)
        assert len(exps) == 1
        assert (exps[0].current_ma, exps[0].cutoff_v, exps[0].threshold_mah) == (30.0, 2.25, 200)

    def test_empty_required_capacity_yields_no_coin_expectation(self):
        # Coin：期望直接取规则/UI 的绝对值 → 空值则无期望格
        ti = make_test_info("CR2032", "", "1S1P", 220, "", [30, 26, 15, 6], VOLTAGES)
        assert build_expectations(ti) == []

    def test_pouch_expectation_derives_from_qnom_even_without_required(self):
        # Pouch：期望按 PDF 由 QNOM × 构造比例派生，不依赖 UI 期望字段
        ti = make_test_info("XYZ1430", "Spiral Type", "1S1P", 1500, "", CURRENTS, VOLTAGES)
        exps = build_expectations(ti)
        assert len(exps) == 1 and exps[0].threshold_mah == 1275


# ── 5. 全链路：统计 → 报告 Pass/Fail ──


class TestFullReport:
    def test_pass_when_all_samples_and_mu2sigma_above_expectation(self):
        ti = make_test_info("XYZ1430", "Spiral Type", "1S1P", 1500, 1275, CURRENTS, VOLTAGES)
        bi = make_battery_info(make_charge([1300, 1310, 1290]))
        out = run_report(ti, bi)
        assert out["strResult"] == "Pass"
        assert out["strRemarks"] == "OK"

    def test_fail_on_single_sample_below_expectation(self):
        ti = make_test_info("XYZ1430", "Spiral Type", "1S1P", 1500, 1275, CURRENTS, VOLTAGES)
        bi = make_battery_info(make_charge([1260, 1300, 1300]))  # 1260 < 1275
        out = run_report(ti, bi)
        assert out["strResult"] == "Fail"
        assert (
            "The minimum sample capacity should be at least 1275mAh" in out["strRemarks"]
        )
        assert "actual measured minimum capacity to 2.25V is 1260mAh" in out["strRemarks"]

    def test_fail_on_mu2sigma_not_higher_than_expectation(self):
        # min ≥ 1275 但离散度大 → µ-2σ < 1275 → 仅统计判据失败
        ti = make_test_info("XYZ1430", "Spiral Type", "1S1P", 1500, 1275, CURRENTS, VOLTAGES)
        bi = make_battery_info(make_charge([1280, 1280, 1600]))
        out = run_report(ti, bi)
        assert out["strResult"] == "Fail"
        assert "The expected usable Q(stat) should be more than 85%" in out["strRemarks"]

    def test_coin_new_model_full_report(self):
        ti = make_test_info("CR2032", "", "1S1P", 220, 200, [30, 26, 15, 6], VOLTAGES)
        # mean=215, 样本σ=5 → µ-2σ=205 > 200；min=210 ≥ 200 → Pass
        bi = make_battery_info(make_charge([215, 220, 210]))
        out = run_report(ti, bi)
        assert out["strResult"] == "Pass"
        assert out["strRemarks"] == "OK"

    def test_mu2sigma_equal_to_threshold_is_fail_strict(self):
        """µ-2σ 恰等于期望（200）不满足 PDF 的 higher than → 严格判 Fail"""
        ti = make_test_info("CR2032", "", "1S1P", 220, 200, [30, 26, 15, 6], VOLTAGES)
        # mean=210, 样本σ=5 → µ-2σ=200.0 == 200 → stat_ok False（单样本 205 ≥ 200 通过）
        bi = make_battery_info(make_charge([210, 215, 205]))
        out = run_report(ti, bi)
        assert out["strResult"] == "Fail"
        assert out["strRemarks"] == (
            "The expected usable Q(stat) should be more than 90%, "
            "while the actual measured minimum capacity to 2.25V is 90%."
        )

    def test_falls_back_to_legacy_when_expected_level_not_measured(self):
        """期望格（60mA）不在实测档位 → 回退历史单点判定（最高电流 × 2.25V）"""
        levels = [30, 26, 15, 6]
        ti = make_test_info("XYZ1430", "Spiral Type", "1S1P", 1500, 1275, levels, VOLTAGES)
        # 回退路径：mm2s[0][4] / qnom ≥ required / qnom → 1300 ≥ 1275 Pass
        bi = make_battery_info(make_charge([1300, 1300, 1300]))
        out = run_report(ti, bi)
        assert out["strResult"] == "Pass"
        # 判定不变；缺少期望档位时必须明确说明，而不是写 OK。
        assert "no measured samples at 60mA → 2.25V" in out["strRemarks"]
        assert "OK" not in out["strRemarks"]

    def test_legacy_fallback_fail_keeps_historical_remarks(self):
        levels = [30, 26, 15, 6]
        ti = make_test_info("XYZ1430", "Spiral Type", "1S1P", 1500, 1275, levels, VOLTAGES)
        bi = make_battery_info(make_charge([1100, 1100, 1100]))
        out = run_report(ti, bi)
        assert out["strResult"] == "Fail"
        assert "The expected usable Q(stat) should be more than 85%" in out["strRemarks"]


# ── 6. Excel 报告落盘（判定标注）──


class TestExcelEndToEnd:
    def test_write_report_annotates_verdict(self, tmp_path):
        ti = make_test_info("XYZ1430", "Spiral Type", "1S1P", 1500, 1275, CURRENTS, VOLTAGES)
        bi = make_battery_info(make_charge([1260, 1300, 1300]))  # 单样本低于期望 → Fail
        (tmp_path / "results").mkdir()  # 真实管线由上游创建结果目录
        writer = ExcelReportWriter(str(tmp_path / "results"), ti, bi)
        writer.write()

        wb = load_workbook(writer.strResultXlsxPath)
        ws = wb["overview"]
        cells = [
            str(cell.value) for row in ws.iter_rows() for cell in row if cell.value is not None
        ]
        annotated = [t for t in cells if "Verdict:" in t]
        assert annotated, "Overview 判定格未标注"
        assert any("Expected: 1275mAh" in t for t in annotated)
        assert any("60mA → 2.25V" in t for t in annotated)
        assert any("Verdict: Fail" in t for t in annotated)
        wb.close()

        # 样本 Excel（sample 内容）里 Result 列应为 Fail
        wb2 = load_workbook(writer.strSampleXlsxPath)
        all_values = [
            str(cell.value)
            for sheet in wb2.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        ]
        assert "Fail" in all_values
        wb2.close()

    def test_write_report_pass_verdict(self, tmp_path):
        ti = make_test_info("XYZ1430", "Spiral Type", "1S1P", 1500, 1275, CURRENTS, VOLTAGES)
        bi = make_battery_info(make_charge([1300, 1310, 1290]))
        (tmp_path / "results").mkdir()
        writer = ExcelReportWriter(str(tmp_path / "results"), ti, bi)
        writer.write()

        wb = load_workbook(writer.strResultXlsxPath)
        cells = [
            str(cell.value)
            for row in wb["overview"].iter_rows()
            for cell in row
            if cell.value is not None
        ]
        assert any("Verdict: Pass" in t for t in cells)
        assert any("Expected: 1275mAh" in t for t in cells)
        wb.close()
